"""Build the testing workbook from the pytest JUnit file and the smoke checks."""

import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

ROOT = Path(__file__).resolve().parents[1]
JUNIT = ROOT / "reports" / "junit.xml"
OUT = ROOT / "docs" / "testing-report.xlsx"

INK = "202B24"
SAGE = "A6B59A"
PAPER = "F4F0E5"
CLAY = "B9785B"
WHITE = "FFFFFF"

header_fill = PatternFill("solid", fgColor=INK)
header_font = Font(name="Arial", bold=True, color=WHITE, size=11)
body_font = Font(name="Arial", size=11)
title_font = Font(name="Arial", bold=True, size=16, color=INK)
note_font = Font(name="Arial", size=10, italic=True, color="5E6057")
pass_fill = PatternFill("solid", fgColor="E4EBD8")
fail_fill = PatternFill("solid", fgColor="F3D9CF")


def cases():
    tree = ET.parse(JUNIT)
    rows = []
    for case in tree.iter("testcase"):
        failed = case.find("failure") is not None or case.find("error") is not None
        rows.append(
            (
                case.get("classname", ""),
                case.get("name", ""),
                "Fail" if failed else "Pass",
                float(case.get("time") or 0),
            )
        )
    return rows


def style_header(sheet, columns):
    for col in range(1, columns + 1):
        cell = sheet.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 22
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddHeader.left.text = "Moss Garden"
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.print_title_rows = "1:1"


def main():
    rows = cases()
    smoke = [
        ("Health", "http://localhost:8000/healthz", 200, "Pass"),
        ("Favicon SVG on API", "http://localhost:8000/favicon.svg", 200, "Pass"),
        ("Favicon ICO on API", "http://localhost:8000/favicon.ico", 200, "Pass"),
        ("Favicon PNG on API", "http://localhost:8000/favicon-32.png", 200, "Pass"),
        ("Apple touch icon on API", "http://localhost:8000/apple-touch-icon.png", 200, "Pass"),
        ("Auth status", "http://localhost:8000/api/auth/status", 200, "Pass"),
        ("Vite home", "http://localhost:5173/", 200, "Pass"),
        ("Favicon SVG on Vite", "http://localhost:5173/favicon.svg", 200, "Pass"),
        ("Demo enter", "POST http://localhost:5173/api/demo/enter", 200, "Pass"),
        ("Demo garden", "GET http://localhost:5173/api/garden", 200, "Pass"),
        ("Logout", "POST http://localhost:5173/api/logout", 200, "Pass"),
        ("Garden after logout", "GET http://localhost:5173/api/garden", 401, "Pass"),
        ("Home links the favicon", "http://localhost:5173/ index HTML", 200, "Pass"),
    ]

    book = Workbook()
    summary = book.active
    summary.title = "Summary"

    summary["A1"] = "Moss Garden test report"
    summary["A1"].font = title_font
    summary.merge_cells("A1:B1")
    summary["A2"] = "Pytest cases come from reports/junit.xml. Totals are formulas over the Cases sheet. Smoke rows are the live checks run on 22 Oct 2026 against localhost."
    summary["A2"].font = note_font
    summary.merge_cells("A2:D2")
    summary.row_dimensions[2].height = 32

    labels = [
        (4, "Automated cases", "=COUNTA(Cases!C2:C2000)"),
        (5, "Passed", '=COUNTIF(Cases!C2:C2000,"Pass")'),
        (6, "Failed", '=COUNTIF(Cases!C2:C2000,"Fail")'),
        (7, "Pass rate", "=IF(B4=0,0,B5/B4)"),
        (8, "Seconds", "=SUM(Cases!D2:D2000)"),
        (10, "Smoke checks", "=COUNTA(Smoke!D2:D50)"),
        (11, "Smoke passed", '=COUNTIF(Smoke!D2:D50,"Pass")'),
        (12, "Smoke failed", '=COUNTIF(Smoke!D2:D50,"Fail")'),
    ]
    for row, label, formula in labels:
        summary.cell(row, 1, label).font = body_font
        cell = summary.cell(row, 2, formula)
        cell.font = Font(name="Arial", size=12, bold=True)
        cell.alignment = Alignment(horizontal="right")
    summary["B7"].number_format = "0.0%"
    summary["B8"].number_format = "0.00"
    summary["A14"] = "A case is one pytest node, including each parameter of a parametrized test. Source: local pytest run, not an external dataset."
    summary["A14"].font = note_font
    summary.merge_cells("A14:D14")

    chart = BarChart()
    chart.type = "col"
    chart.title = "Automated results"
    chart.y_axis.title = "Cases"
    chart.style = 10
    summary["A16"] = "Result"
    summary["B16"] = "Count"
    summary["A16"].font = header_font
    summary["B16"].font = header_font
    summary["A16"].fill = header_fill
    summary["B16"].fill = header_fill
    summary["A17"] = "Passed"
    summary["B17"] = "=B5"
    summary["A18"] = "Failed"
    summary["B18"] = "=B6"
    data = Reference(summary, min_col=2, min_row=16, max_row=18)
    cats = Reference(summary, min_col=1, min_row=17, max_row=18)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.y_axis.scaling.min = 0
    summary.add_chart(chart, "D4")

    sheet = book.create_sheet("Cases")
    for col, name in enumerate(("File", "Case", "Result", "Seconds"), 1):
        sheet.cell(1, col, name)
    for index, (classname, name, outcome, seconds) in enumerate(rows, 2):
        sheet.cell(index, 1, classname).font = body_font
        sheet.cell(index, 2, name).font = body_font
        result = sheet.cell(index, 3, outcome)
        result.font = body_font
        result.fill = pass_fill if outcome == "Pass" else fail_fill
        timing = sheet.cell(index, 4, seconds)
        timing.font = body_font
        timing.number_format = "0.000"
    style_header(sheet, 4)
    sheet.auto_filter.ref = f"A1:D{len(rows) + 1}"
    widths = (42, 72, 12, 12)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = f"A1:D{max(len(rows) + 1, 2)}"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_TABLOID

    smoke_sheet = book.create_sheet("Smoke")
    for col, name in enumerate(("Check", "Target", "Status", "Result"), 1):
        smoke_sheet.cell(1, col, name)
    for index, item in enumerate(smoke, 2):
        for col, value in enumerate(item, 1):
            cell = smoke_sheet.cell(index, col, value)
            cell.font = body_font
            if col == 4:
                cell.fill = pass_fill if value == "Pass" else fail_fill
    style_header(smoke_sheet, 4)
    smoke_sheet.auto_filter.ref = f"A1:D{len(smoke) + 1}"
    for index, width in enumerate((28, 56, 12, 12), 1):
        smoke_sheet.column_dimensions[get_column_letter(index)].width = width

    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 16
    summary.column_dimensions["C"].width = 18
    summary.column_dimensions["D"].width = 22
    summary.page_setup.orientation = "landscape"
    summary.page_setup.fitToPage = True
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 1
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.oddFooter.right.text = "Page &P of &N"
    summary.oddHeader.left.text = "Moss Garden test report"
    summary.freeze_panes = "A4"
    summary.sheet_view.showGridLines = False
    summary.sheet_view.zoomScale = 120
    summary.print_title_rows = "1:3"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUT)
    print(f"cases {len(rows)} smoke {len(smoke)} -> {OUT}")


if __name__ == "__main__":
    main()
